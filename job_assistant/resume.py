from pathlib import Path
from docx import Document

def read_resume(path: str) -> str:
    p = Path(path)
    if not p.exists():
        raise FileNotFoundError(f"Resume not found: {p}. Add your base resume DOCX first.")
    doc = Document(p)
    return "\n".join(p.text for p in doc.paragraphs if p.text.strip())

def save_tailored_docx(text: str, output_path: str):
    doc = Document()
    for line in text.splitlines():
        doc.add_paragraph(line)
    Path(output_path).parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)
