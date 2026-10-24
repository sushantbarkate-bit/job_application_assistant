import sqlite3
import tempfile
import unittest
from datetime import datetime, timedelta
from pathlib import Path
from unittest import mock

from docx import Document

from job_assistant import cli
from job_assistant.config import ROOT, load_config
from job_assistant.matcher import score_job
from job_assistant.models import Job
from job_assistant.providers import GreenhouseProvider, clean_description
from job_assistant.report import daily_report
from job_assistant.storage import Store


def greenhouse_item(job_id, title, content="", location="Pune, India"):
    return {
        "id": job_id,
        "title": title,
        "content": content,
        "location": {"name": location},
        "absolute_url": f"https://boards.greenhouse.io/acme/jobs/{job_id}",
    }


class GreenhouseMatchingTests(unittest.TestCase):
    def search_with(self, items, keywords, locations):
        provider = GreenhouseProvider(["acme"])
        with mock.patch.object(provider, "_fetch_board", return_value=items):
            return provider.search(keywords, locations, 50)

    def test_engineer_title_matches_developer_keyword(self):
        jobs = self.search_with(
            [greenhouse_item(1, "Frontend Engineer")],
            ["Frontend Developer"], ["Pune"],
        )
        self.assertEqual([job.title for job in jobs], ["Frontend Engineer"])

    def test_keyword_phrase_in_html_description_matches(self):
        content = "&lt;p&gt;We need a &lt;b&gt;Java Full Stack Developer&lt;/b&gt;.&lt;/p&gt;"
        jobs = self.search_with(
            [greenhouse_item(2, "Software Engineer III", content)],
            ["Java Full Stack Developer"], ["Pune"],
        )
        self.assertEqual(len(jobs), 1)
        self.assertIn("Java Full Stack Developer", jobs[0].description)
        self.assertNotIn("<", jobs[0].description)

    def test_location_alias_bangalore_matches_bengaluru(self):
        jobs = self.search_with(
            [greenhouse_item(3, "Angular Developer", location="Bangalore, India")],
            ["Angular Developer"], ["Bengaluru"],
        )
        self.assertEqual(len(jobs), 1)

    def test_unrelated_titles_and_locations_are_excluded(self):
        jobs = self.search_with(
            [
                greenhouse_item(4, "Data Scientist"),
                greenhouse_item(5, "Angular Developer", location="New York, NY"),
            ],
            ["Angular Developer"], ["Pune"],
        )
        self.assertEqual(jobs, [])

    def test_clean_description_strips_escaped_html(self):
        self.assertEqual(clean_description("&lt;p&gt;Hello&amp;nbsp;&lt;/p&gt;"), "Hello")


class MatcherTests(unittest.TestCase):
    def test_java_does_not_match_inside_javascript(self):
        job = Job("m-1", "Frontend Developer", "Co", "Pune",
                  "Strong JavaScript and TypeScript skills", "https://example.com")
        result = score_job(job, {"skills": ["Java", "JavaScript"]})
        self.assertEqual(result["matched_skills"], ["javascript"])

    def test_engineer_title_counts_as_developer_target(self):
        job = Job("m-3", "Senior Frontend Engineer", "Co", "Bangalore, India",
                  "Angular and TypeScript", "https://example.com")
        result = score_job(job, {"skills": ["Angular", "TypeScript"],
                                 "target_titles": ["Frontend Developer"],
                                 "target_locations": ["Bengaluru"]})
        self.assertTrue(result["title_match"])
        self.assertTrue(result["location_match"])

    def test_skill_at_word_boundary_still_matches(self):
        job = Job("m-2", "Developer", "Co", "Pune", "Java, Spring Boot.", "https://example.com")
        result = score_job(job, {"skills": ["Java", "Spring Boot"]})
        self.assertEqual(result["matched_skills"], ["java", "spring boot"])


class StorageTests(unittest.TestCase):
    def make_legacy_db(self, path):
        db = sqlite3.connect(path)
        db.execute("""CREATE TABLE applications (
            job_id TEXT PRIMARY KEY, title TEXT, company TEXT, location TEXT, url TEXT,
            score INTEGER, status TEXT, created_at TEXT, notes TEXT)""")
        old = (datetime.now() - timedelta(days=3)).isoformat(timespec="seconds")
        db.execute("INSERT INTO applications VALUES (?,?,?,?,?,?,?,?,?)",
                   ("old-1", "Angular Developer", "Acme", "Pune", "https://example.com/1",
                    60, "matched", old, "Matched: angular"))
        db.commit()
        db.close()

    def test_legacy_database_is_migrated_without_losing_rows(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = str(Path(tmp) / "legacy.sqlite3")
            self.make_legacy_db(path)
            store = Store(path)
            columns = {row[1] for row in store.db.execute("PRAGMA table_info(applications)")}
            self.assertIn("description", columns)
            self.assertIn("updated_at", columns)
            self.assertEqual(store.rows()[0][0], "old-1")
            store.db.close()

    def test_job_description_is_stored_and_backfilled_on_rediscovery(self):
        with tempfile.TemporaryDirectory() as tmp:
            store = Store(str(Path(tmp) / "jobs.sqlite3"))
            store.add(Job("b-1", "Angular Developer", "Acme", "Pune", "", "https://example.com"), 60)
            self.assertEqual(store.job("b-1")["description"], "")
            store.add(Job("b-1", "Angular Developer", "Acme", "Pune", "Build Angular", "https://example.com"), 60)
            self.assertEqual(store.job("b-1")["description"], "Build Angular")
            store.db.close()

    def test_rediscovery_refreshes_unacted_jobs_but_keeps_approved_ones(self):
        with tempfile.TemporaryDirectory() as tmp:
            store = Store(str(Path(tmp) / "jobs.sqlite3"))
            job = Job("r-1", "Frontend Engineer", "Co", "Pune", "Angular", "https://example.com")
            store.add(job, 30, "low_match", "old note")
            store.add(job, 70, "matched", "new note")
            self.assertEqual(store.rows()[0][5:7], (70, "matched"))

            store.update_status("r-1", "approved_for_review", "User approved review")
            store.add(job, 10, "low_match", "ignored")
            self.assertEqual(store.rows()[0][5:7], (70, "approved_for_review"))
            store.db.close()

    def test_approved_job_appears_in_report_on_a_later_day(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = str(Path(tmp) / "jobs.sqlite3")
            self.make_legacy_db(path)
            store = Store(path)
            store.update_status("old-1", "approved_for_review", "User approved")
            report_path, counts = daily_report(store, str(Path(tmp) / "reports"))
            self.assertEqual(counts, {"approved_for_review": 1})
            self.assertIn("old-1", report_path.read_text(encoding="utf-8-sig"))
            store.db.close()


class ConfigAndLinksTests(unittest.TestCase):
    def test_relative_paths_resolve_against_project_root(self):
        cfg = load_config()
        self.assertEqual(Path(cfg["workflow"]["database_path"]), ROOT / "data" / "applications.sqlite3")
        self.assertEqual(Path(cfg["workflow"]["output_dir"]), ROOT / "output")
        self.assertTrue(Path(cfg["profile"]["resume_path"]).is_absolute())

    def test_disabled_sources_are_not_shown(self):
        cfg = {
            "search": {
                "keywords": ["Angular Developer"],
                "sources": [
                    {"name": "LinkedIn", "url": "https://www.linkedin.com/jobs/", "enabled": True},
                    {"name": "Naukri", "url": "https://www.naukri.com/", "enabled": False},
                ],
            },
            "profile": {"target_locations": ["Pune"]},
        }
        links = cli.build_portal_links(cfg)
        self.assertEqual({link["source"] for link in links}, {"LinkedIn"})

    def test_all_sources_shown_when_none_configured(self):
        cfg = {"search": {"keywords": ["Angular Developer"]}, "profile": {"target_locations": ["Pune"]}}
        self.assertEqual({link["source"] for link in cli.build_portal_links(cfg)}, {"LinkedIn", "Naukri"})


class TailorResumeTests(unittest.TestCase):
    def test_tailored_resume_is_saved_as_docx(self):
        with tempfile.TemporaryDirectory() as tmp:
            store = Store(str(Path(tmp) / "jobs.sqlite3"))
            store.add(Job("t-1", "Angular Developer", "Acme", "Pune",
                          "Build Angular components", "https://example.com"), 60)
            resume = Path(tmp) / "resume.docx"
            document = Document()
            document.add_paragraph("Sushant Barkate")
            document.save(resume)
            cfg = {"profile": {"resume_path": str(resume)},
                   "workflow": {"output_dir": str(Path(tmp) / "output")}}

            with mock.patch("job_assistant.cli.AIClient.tailor_resume_bullets",
                            return_value="Tailored line [USER VERIFY]"):
                cli.tailor_resume(cfg, store, "t-1")

            saved = Path(tmp) / "output" / "tailored-resumes" / "t-1.docx"
            self.assertTrue(saved.exists())
            self.assertIn("Tailored line [USER VERIFY]",
                          "\n".join(p.text for p in Document(saved).paragraphs))
            self.assertEqual(store.job("t-1")["status"], "matched")
            store.db.close()

    def test_unknown_job_is_refused_without_calling_ai(self):
        with tempfile.TemporaryDirectory() as tmp:
            store = Store(str(Path(tmp) / "jobs.sqlite3"))
            with mock.patch("job_assistant.cli.AIClient.tailor_resume_bullets") as ai:
                cli.tailor_resume({"profile": {}, "workflow": {}}, store, "missing")
            ai.assert_not_called()
            store.db.close()


if __name__ == "__main__":
    unittest.main()
